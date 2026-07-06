import argparse
import csv
import math
import re
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook

# Sheet mapping: sheet_name -> (sku_column, product_name_column, price_column, extra_price_multiplier)
SHEET_MAPPINGS: Dict[str, Tuple[str, str, str, float]] = {
    "COUGAR GAMING": ("B", "C", "G", 1.0),
    "COUGAR ACC": ("B", "C", "G", 1.0),
    "COUGAR PC": ("B", "C", "F", 1.0),
    "ANTEC": ("B", "C", "G", 1.0),
    "ANTEC PSU": ("B", "C", "F", 1.0),
    "ANTEC COOL": ("B", "C", "G", 1.0),
    "NOCTUA": ("B", "C", "F", 3.0),
    "CREATIVE": ("B", "C", "H", 1.0),
    "BE QUIET": ("B", "C", "G", 1.0),
    "LEADTEC GPU": ("A", "C", "G", 1.0),
}


def parse_numeric_price(value: object) -> Optional[float]:
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("\u00a0", "")
    cleaned = re.sub(r"[^0-9,.-]", "", text)
    if not cleaned:
        return None

    if "," in cleaned and "." in cleaned:
        cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        parts = cleaned.split(",")
        if len(parts) == 2 and len(parts[1]) <= 2:
            cleaned = cleaned.replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")

    try:
        return float(cleaned)
    except ValueError:
        return None


def round_up_to_next_9(value: float) -> int:
    rounded_up = math.ceil(value)
    remainder = rounded_up % 10
    if remainder == 9:
        return rounded_up
    return rounded_up + (9 - remainder)


def calculate_final_price(base_price: float) -> int:
    calculated = base_price * 1.18 * 1.3
    return round_up_to_next_9(calculated)


def read_rows(
    xlsx_path: str,
    supplier_name: str,
) -> Iterable[List[object]]:
    workbook = load_workbook(filename=xlsx_path, data_only=True, read_only=True)

    for sheet_name, mapping in SHEET_MAPPINGS.items():
        if sheet_name not in workbook.sheetnames:
            print(f"[WARN] Sheet not found: {sheet_name}")
            continue

        sku_col, product_col, price_col, multiplier = mapping
        sheet = workbook[sheet_name]

        for row_idx in range(1, sheet.max_row + 1):
            supplier_sku = sheet[f"{sku_col}{row_idx}"].value
            product_name = sheet[f"{product_col}{row_idx}"].value
            raw_price = sheet[f"{price_col}{row_idx}"].value

            if supplier_sku is None or product_name is None:
                continue

            price = parse_numeric_price(raw_price)
            if price is None or price <= 0:
                continue

            price *= multiplier
            final_price = calculate_final_price(price)

            yield [
                supplier_name,
                str(supplier_sku).strip(),
                str(product_name).strip(),
                "",
                final_price,
            ]


def write_output(rows: Iterable[List[object]], output_path: str, with_header: bool) -> int:
    count = 0
    with open(output_path, "w", newline="", encoding="utf-8-sig") as output_file:
        writer = csv.writer(output_file)

        if with_header:
            writer.writerow(["supplier_name", "supplier_sku", "product_name", "empty_field", "calculated_price"])

        for row in rows:
            if not row[1] or not row[2]:
                continue
            writer.writerow(row)
            count += 1

    return count


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Parse Five supplier XLSX and export supplier import lines.",
    )
    parser.add_argument("xlsx_path", help="Path to the Five supplier XLSX file")
    parser.add_argument(
        "--supplier-name",
        required=True,
        help="Supplier name to place in every output row (selected from form)",
    )
    parser.add_argument(
        "--output",
        default="five_supplier_import.csv",
        help="Output CSV file path (default: five_supplier_import.csv)",
    )
    parser.add_argument(
        "--with-header",
        action="store_true",
        help="Include CSV header row",
    )
    return parser


def main() -> None:
    args = build_parser().parse_args()
    rows = read_rows(args.xlsx_path, args.supplier_name)
    written = write_output(rows, args.output, with_header=args.with_header)
    print(f"Done. Wrote {written} rows to {args.output}")


if __name__ == "__main__":
    main()
